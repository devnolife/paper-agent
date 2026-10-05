"""Deterministic tests (no network, no models)."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from paper_agent import agent, citations, docx2tex, latex_mask
from paper_agent.latex_mask import Doc, fix_ref_words, integrity_ok, mask_text, unmask

SAMPLE_TEX = r"""\section{System Design}
\label{sec:method}

Fig.~\ref{fig:arch} shows the four phases~\cite{hevner2004}. Verdicts are \textsc{pass} or \textsc{flag}; 1{,}735~s---98.7\% of the run.

\begin{table}[t]
  \caption{The nine validation rules.}
  \label{tab:rules}
  \centering
  \begin{tabular}{@{}llp{4.3cm}@{}}
    \toprule
    ID & Family & Condition $\rightarrow$ action \\
    \midrule
    F1 & Feasibility & method needs high resources $\wedge$ low-resource $\rightarrow$ \textsc{reject} \\
    \bottomrule
  \end{tabular}
\end{table}

\paragraph{Fragmentation}
Papers address the same phenomenon (\emph{semantic} filtering keeps pairs with $Q\le0.42$); see Section~\ref{sec:rules}. Only 58\% recurred.
\begin{itemize}
  \item \textbf{RQ1} To what extent can an agentic system detect indicators?
\end{itemize}
"""


def test_mask_roundtrip_is_identity(tmp_path: Path):
    (tmp_path / "method.tex").write_text(SAMPLE_TEX, encoding="utf-8")
    doc = Doc()
    doc.add_file("method", SAMPLE_TEX)
    assert doc.render("method") == SAMPLE_TEX
    assert len(doc.units) >= 8
    texts = [u.masked for u in doc.units]
    assert any("<1>semantic</1>" in t for t in texts)
    assert all("\\cite" not in t and "$" not in t for t in texts)


def test_integrity_rejects_moved_or_missing_placeholders():
    masked, ph, tags = mask_text(r"Only 58\% recurred in \emph{two} runs~\cite{a}.")
    assert integrity_ok(masked, masked)
    moved = masked.replace("58" + ph_token(masked, 0), ph_token(masked, 0) + "58")
    assert not integrity_ok(masked, moved)
    assert not integrity_ok(masked, masked.replace("<1>", "").replace("</1>", ""))
    assert not integrity_ok(masked, masked + ph_token(masked, 0))


def ph_token(masked: str, idx: int) -> str:
    return f"{latex_mask.PH_OPEN}{idx}{latex_mask.PH_CLOSE}"


def test_unmask_restores_latex_and_escapes_specials():
    masked, ph, tags = mask_text(r"Rate 58\% with \textbf{bold} and \texttt{nli}.")
    translated = masked.replace("Rate", "Laju").replace("bold", "tebal") + " 100% & more"
    out = unmask(translated, ph, tags)
    assert out.startswith("Laju 58\\% with \\textbf{tebal} and \\texttt{nli}.")
    assert out.endswith("100\\% \\& more")
    assert unmask("\u201cquoted\u201d \u2014 dash", [], {}) == "``quoted'' --- dash"


def test_fix_ref_words():
    assert fix_ref_words(r"see Section~\ref{sec:a} and Fig.~\ref{fig:b}") == r"see Bagian~\ref{sec:a} and Gambar~\ref{fig:b}"


def test_citation_gate(tmp_path: Path):
    (tmp_path / "main.tex").write_text(r"\input{sections/a} \cite{x2020}", encoding="utf-8")
    (tmp_path / "sections").mkdir()
    (tmp_path / "sections" / "a.tex").write_text(r"Text~\cite{y2021,z2022}. % \cite{commented}", encoding="utf-8")
    (tmp_path / "refs.bib").write_text("@article{x2020, title={X}}\n@misc{y2021, title={Y}}\n@misc{unused, title={U}}\n", encoding="utf-8")
    (tmp_path / "report.md").write_text("| key | status | sources | DOI | sim | notes |\n|---|---|---|---|---:|---|\n"
                                        "| x2020 | VERIFIED | a, b | - | 1.00 | |\n| y2021 | SINGLE_SOURCE | a | - | 0.9 | |\n", encoding="utf-8")
    res = citations.check_citations(tmp_path / "main.tex", tmp_path / "refs.bib", tmp_path / "report.md")
    assert not res["ok"]
    whys = {p["key"]: p["why"] for p in res["problems"]}
    assert whys == {"y2021": "status SINGLE_SOURCE", "z2022": "missing from .bib"}
    assert res["unused"] == ["unused"]
    assert "commented" not in whys
    assert citations.check_citations(tmp_path / "main.tex", tmp_path / "refs.bib", tmp_path / "report.md", allow_single=True)["problems"] == [
        {"key": "z2022", "why": "missing from .bib", "in": ["a.tex"]}]


def test_docx2tex_produces_compilable_skeleton(tmp_path: Path):
    d = Document()
    d.add_paragraph("A Study of Things", style="Title")
    d.add_paragraph("First Author, Second Author")
    d.add_paragraph("Abstract—We study things with 95% accuracy & care.")
    d.add_paragraph("Index Terms—things, study")
    d.add_paragraph("I. INTRODUCTION")
    p = d.add_paragraph("Things matter [1], [2]–[3]. ")
    p.add_run("Bold claim").bold = True
    d.add_paragraph("A. Background")
    d.add_paragraph("More text_with underscores.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Model", "RMSE", "RF", "0.5"
    d.add_paragraph("REFERENCES")
    d.add_paragraph("[1] A. Author, “Paper one,” Journal, 2020.")
    d.add_paragraph("[2] B. Author, Paper two, 2021.")
    d.add_paragraph("[3] C. Author, Paper three, 2022.")
    src = tmp_path / "in.docx"
    d.save(str(src))
    stats = docx2tex.docx2tex(src, tmp_path / "out")
    body = (tmp_path / "out" / "sections" / "body.tex").read_text(encoding="utf-8")
    main = (tmp_path / "out" / "main.tex").read_text(encoding="utf-8")
    bib = (tmp_path / "out" / "references.bib").read_text(encoding="utf-8")
    assert stats["headings"] == 2 and stats["tables"] == 1 and stats["references"] == 3
    assert "\\title{A Study of Things}" in main
    assert "\\begin{abstract}" in body and "95\\% accuracy \\& care" in body
    assert "\\section{Introduction}" in body and "\\subsection{Background}" in body
    assert "\\cite{ref1}" in body and "\\cite{ref2,ref3}" in body and "\\textbf{Bold claim}" in body
    assert "text\\_with" in body and "\\begin{tabular}" in body
    assert bib.count("@misc{") == 3


def test_keyword_router_builds_plan(tmp_path: Path):
    (tmp_path / "paper" / "sections").mkdir(parents=True)
    (tmp_path / "paper" / "main.tex").write_text("x", encoding="utf-8")
    (tmp_path / "paper" / "main_id.tex").write_text("x", encoding="utf-8")
    plan = agent.plan_by_keywords("terjemahkan paper ke bahasa Indonesia lalu kompilasi", tmp_path)
    assert [s["tool"] for s in plan["steps"]] == ["translate_tex", "compile"]
    assert plan["steps"][0]["args"]["src_dir"] == "paper/sections"
    with pytest.raises(ValueError):
        agent._path(tmp_path, "../outside.tex")


def test_execute_reports_unknown_tool(tmp_path: Path):
    res = agent.execute({"steps": [{"tool": "nope", "args": {}}]}, tmp_path, log=lambda *_: None)
    assert res["ok"] is False and "unknown tool" in res["results"][0]["result"]["error"]
    assert json.dumps(res)  # serialisable
